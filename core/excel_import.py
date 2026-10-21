import re
from openpyxl import load_workbook

USERNAME_RX = re.compile(r'^[\w.@+-]+$')

# Başlıq adları (fərqli yazılışlar qəbul olunur)
HEADERS = {
    'first_name': {'ad', 'first_name', 'name'},
    'last_name': {'soyad', 'last_name', 'surname'},
    'username': {'istifadəçi adı', 'istifadeci adi', 'istifadəçi', 'username', 'login'},
    'password': {'şifrə', 'sifre', 'şifre', 'password'},
    'group': {'qrup', 'group'},
}


def _norm(s):
    """Azərbaycan hərflərinə uyğun kiçik hərfə çevirmə (İ, I, ı -> i)."""
    return s.replace('İ', 'i').replace('I', 'i').replace('ı', 'i').lower().strip()


def _clean(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_students_excel(file, group_names, existing_users, institution_id):
    """
    Excel-i oxuyur, bazaya YAZMIR.
    group_names: müəssisənin mövcud qrup adları (set)
    existing_users: {username: (role, institution_id)}
    Qaytarır: (rows, fatal_error)
    """
    try:
        wb = load_workbook(file, read_only=True, data_only=True)
        ws = wb.active
        data = list(ws.iter_rows(values_only=True))
    except Exception:
        return [], "Excel faylını oxumaq olmadı. .xlsx formatında olduğuna əmin ol."

    if not data:
        return [], "Fayl boşdur."

    header = [_norm(_clean(c)) for c in data[0]]
    col = {}
    for key, names in HEADERS.items():
        for i, h in enumerate(header):
            if h in {_norm(x) for x in names}:
                col[key] = i
                break
    missing = [k for k in ('first_name', 'last_name', 'username', 'password', 'group') if k not in col]
    if missing:
        return [], "Başlıq sətrində bu sütunlar tapılmadı: " + ", ".join(
            {'first_name': 'Ad', 'last_name': 'Soyad', 'username': 'İstifadəçi adı',
             'password': 'Şifrə', 'group': 'Qrup'}[m] for m in missing)

    rows = []
    seen = set()
    lower_groups = {_norm(g): g for g in group_names}

    for n, raw in enumerate(data[1:], start=2):
        if raw is None or all(_clean(c) == '' for c in raw):
            continue  # boş sətir

        def get(key):
            i = col[key]
            return _clean(raw[i]) if i < len(raw) else ''

        row = {
            'row': n,
            'first_name': get('first_name'),
            'last_name': get('last_name'),
            'username': get('username'),
            'password': get('password'),
            'group': get('group'),
            'action': 'create',
            'errors': [],
        }
        errs = row['errors']

        if not row['username']:
            errs.append("İstifadəçi adı boşdur")
        elif not USERNAME_RX.match(row['username']):
            errs.append("İstifadəçi adında boşluq və ya qadağan simvol var")
        elif row['username'].lower() in seen:
            errs.append("Bu istifadəçi adı faylda təkrarlanır")
        seen.add(row['username'].lower())

        if not row['first_name'] or not row['last_name']:
            errs.append("Ad və ya soyad boşdur")

        if not row['group']:
            errs.append("Qrup boşdur")
        elif _norm(row['group']) not in lower_groups:
            errs.append(f"'{row['group']}' qrupu tapılmadı (əvvəl Qruplar səhifəsində yarat)")
        else:
            row['group'] = lower_groups[_norm(row['group'])]

        existing = existing_users.get(row['username'])
        if existing is not None:
            role, inst = existing
            if role == 'student' and inst == institution_id:
                row['action'] = 'add'      # yeni hesab yox, yalnız qrupa əlavə
            else:
                errs.append("Bu istifadəçi adı artıq başqa hesaba məxsusdur")
        else:
            if len(row['password']) < 6:
                errs.append("Şifrə ən azı 6 simvol olmalıdır")

        rows.append(row)

    if not rows:
        return [], "Faylda tələbə sətri tapılmadı."
    return rows, None